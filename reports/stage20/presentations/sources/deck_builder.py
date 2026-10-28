"""
WorkSphere Deck Builder Module
Common visual design system for 16:9 widescreen academic presentations.
Supports professional color palette, structured cards, styled tables, code boxes, images, and comprehensive speaker notes.
"""
import os
import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# 16:9 Widescreen dimensions
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Color Palette
C_NAVY_DARK = RGBColor(0x0A, 0x19, 0x2F)     # Midnight Blue (Cover background)
C_NAVY = RGBColor(0x0D, 0x47, 0xA1)          # Deep Royal Navy (Primary Header)
C_BLUE = RGBColor(0x1E, 0x88, 0xE5)          # Bright Blue (Accents, active borders)
C_TEAL = RGBColor(0x00, 0x97, 0xA7)          # Cyan/Teal (Secondary highlights)
C_DARK = RGBColor(0x21, 0x21, 0x21)          # Charcoal (Body text)
C_MUTED = RGBColor(0x75, 0x75, 0x75)         # Muted Gray
C_BG_CARD = RGBColor(0xF5, 0xF9, 0xFD)       # Light Blue card background
C_BG_WHITE = RGBColor(0xFF, 0xFF, 0xFF)      # Pure White
C_BORDER = RGBColor(0xCF, 0xD8, 0xDC)        # Light Slate border
C_RED = RGBColor(0xD3, 0x2F, 0x2F)           # Red for warnings/critical
C_ORANGE = RGBColor(0xF5, 0x7C, 0x00)        # Orange for concept/partial

def create_presentation():
    prs = pptx.Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    return prs

def add_header(slide, title_text, category_text, course_tag="WORKSPHERE"):
    # Category tag
    tb_cat = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.4))
    tf_cat = tb_cat.text_frame
    tf_cat.word_wrap = True
    p_cat = tf_cat.paragraphs[0]
    p_cat.text = f"{course_tag.upper()} • {category_text.upper()}"
    p_cat.font.name = "Arial"
    p_cat.font.size = Pt(11)
    p_cat.font.bold = True
    p_cat.font.color.rgb = C_BLUE

    # Title
    tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.8))
    tf_title = tb_title.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title_text
    p_title.font.name = "Arial"
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = C_NAVY

def add_footer(slide, course_name, slide_num, total_slides):
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(11.733), Inches(0.4))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = f"Dự án WorkSphere | Học phần: {course_name} — Trang {slide_num}/{total_slides}"
    p.font.name = "Arial"
    p.font.size = Pt(9.5)
    p.font.color.rgb = C_MUTED

def set_speaker_notes(slide, notes_dict):
    """
    Sets structured speaker notes according to the WorkSphere Speaker Notes Contract:
    - Objective
    - Suggested Narration
    - Duration (PROJECT_PROPOSAL)
    - Speaker (PROPOSED_ASSIGNMENT_PENDING_USER_APPROVAL)
    - Evidence
    - Likely Question & Safe Answer
    - Transition / Demo Cue
    """
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    
    sections = [
        f"MỤC TIÊU SLIDE: {notes_dict.get('objective', '')}",
        f"\nLỜI THUYẾT MINH GỢI Ý:\n{notes_dict.get('narration', '')}",
        f"\nTHỜI LƯỢNG DỰ KIẾN: {notes_dict.get('duration', '1 - 2 phút')} (PROJECT_PROPOSAL)",
        f"NGƯỜI THUYẾT TRÌNH: {notes_dict.get('speaker', '')} (PROPOSED_ASSIGNMENT_PENDING_USER_APPROVAL)",
        f"CĂN CỨ MINH CHỨNG: {notes_dict.get('evidence', '')}",
        f"CÂU HỎI HỘI ĐỒNG DỰ KIẾN: {notes_dict.get('question', '')}",
        f"CÂU TRẢ LỜI AN TOÀN: {notes_dict.get('answer', '')}",
        f"GỢI Ý CHUYỂN SLIDE / DEMO: {notes_dict.get('transition', '')}"
    ]
    if 'visual_attribution' in notes_dict:
        sections.append(f"NGUỒN HÌNH ẢNH: {notes_dict['visual_attribution']}")
        
    tf.text = "\n".join(sections)

def add_card(slide, left, top, width, height, title="", items=None, border_color=C_BLUE, bg_color=C_BG_CARD):
    # Shape background
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.5)
    
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.2)
    tf.margin_bottom = Inches(0.2)
    
    # Title
    if title:
        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.name = "Arial"
        p0.font.size = Pt(15)
        p0.font.bold = True
        p0.font.color.rgb = C_NAVY
        p0.space_after = Pt(8)
    
    # Items
    if items:
        start_idx = 0 if not title else 1
        for i, item in enumerate(items):
            p = tf.add_paragraph() if (title or i > 0) else tf.paragraphs[0]
            p.text = f"• {item}"
            p.font.name = "Arial"
            p.font.size = Pt(12)
            p.font.color.rgb = C_DARK
            p.space_after = Pt(4)
            p.line_spacing = 1.15
            
    return shape

def add_cover_slide(prs, title, subtitle, course_name, team_members, notes_dict):
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    
    # Background navy dark rectangle
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_WIDTH, SLIDE_HEIGHT)
    bg.fill.solid()
    bg.fill.fore_color.rgb = C_NAVY_DARK
    bg.line.fill.background()
    
    # Institution
    tb_inst = slide.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.333), Inches(0.6))
    tf_inst = tb_inst.text_frame
    p_inst = tf_inst.paragraphs[0]
    p_inst.text = "TRƯỜNG ĐẠI HỌC CMC • KHOA CÔNG NGHỆ THÔNG TIN & TRUYỀN THÔNG"
    p_inst.font.name = "Arial"
    p_inst.font.size = Pt(13)
    p_inst.font.bold = True
    p_inst.font.color.rgb = C_BLUE
    
    # Course
    tb_course = slide.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.333), Inches(0.6))
    tf_course = tb_course.text_frame
    p_course = tf_course.paragraphs[0]
    p_course.text = f"BÁO CÁO BẢO VỆ HỌC PHẦN: {course_name.upper()}"
    p_course.font.name = "Arial"
    p_course.font.size = Pt(15)
    p_course.font.bold = True
    p_course.font.color.rgb = RGBColor(0x80, 0xD8, 0xFF)
    
    # Title
    tb_title = slide.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(1.6))
    tf_title = tb_title.text_frame
    tf_title.word_wrap = True
    p_title = tf_title.paragraphs[0]
    p_title.text = title
    p_title.font.name = "Arial"
    p_title.font.size = Pt(32)
    p_title.font.bold = True
    p_title.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    
    if subtitle:
        p_sub = tf_title.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.name = "Arial"
        p_sub.font.size = Pt(17)
        p_sub.font.color.rgb = RGBColor(0xCF, 0xD8, 0xDC)
        p_sub.space_before = Pt(8)
        
    # Team members box
    tb_team = slide.shapes.add_textbox(Inches(1.0), Inches(4.5), Inches(11.333), Inches(2.0))
    tf_team = tb_team.text_frame
    p_team_hdr = tf_team.paragraphs[0]
    p_team_hdr.text = "NHÓM SINH VIÊN THỰC HIỆN (Học kỳ Fall 2026):"
    p_team_hdr.font.name = "Arial"
    p_team_hdr.font.size = Pt(12)
    p_team_hdr.font.bold = True
    p_team_hdr.font.color.rgb = RGBColor(0x80, 0xD8, 0xFF)
    
    for m in team_members:
        p_m = tf_team.add_paragraph()
        p_m.text = f"• {m['name']} ({m['id']}) — {m['role']}"
        p_m.font.name = "Arial"
        p_m.font.size = Pt(11.5)
        p_m.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        p_m.space_before = Pt(2)
        
    p_notice = tf_team.add_paragraph()
    p_notice.text = "*(Trạng thái phân công: PROPOSED_ASSIGNMENT_PENDING_USER_APPROVAL — Chờ phê duyệt chính thức)*"
    p_notice.font.name = "Arial"
    p_notice.font.size = Pt(9.5)
    p_notice.font.italic = True
    p_notice.font.color.rgb = RGBColor(0xB0, 0xBE, 0xC5)
    p_notice.space_before = Pt(4)
    
    set_speaker_notes(slide, notes_dict)
    return slide
